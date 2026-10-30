# -*- coding: utf-8 -*-
"""T03659：把 docs/assets/product-shots 下的真实界面截图嵌入产品介绍 docx。

正文每个「【此处插入功能截图：…】」占位符都包在一张单格虚线表里，
本脚本按占位符文案匹配图片、清空单元格、去边框与底色、插入图片与图注。
"""
import atexit
import io
import os
import shutil
import sys
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX = os.path.join(ROOT, "docs", "XianyuHunter-产品介绍.docx")
SHOTS = os.path.join(ROOT, "docs", "assets", "product-shots")

DESKTOP_CM = 12.0
MOBILE_CM = 5.0

# 占位符文案（冒号后的部分） → [(图片文件名, 图注, 宽度 cm)]
PLAN = {
    "任务管理页": [("01-tasks.png", "图 2-1　任务管理页", DESKTOP_CM)],
    "商品列表页": [("02-items.png", "图 2-2　商品库（卡片视图）", DESKTOP_CM)],
    "评估明细页": [("03-evaluations.png", "图 2-3　评估明细与分布", DESKTOP_CM)],
    "价格策略配置页": [("04-price-strategy.png", "图 2-4　价格策略配置", DESKTOP_CM)],
    "订单页 / 购买确认": [
        ("05-orders.png", "图 2-5a　抢单记录（订单状态流转）", DESKTOP_CM),
        ("05b-confirm-buy.png", "图 2-5b　购买确认与商品快照", DESKTOP_CM),
    ],
    "通知中心 / 渠道配置": [
        ("06b-notifications.png", "图 2-6a　通知中心", DESKTOP_CM),
        ("06-notifier.png", "图 2-6b　通知渠道与免打扰", DESKTOP_CM),
    ],
    "移动端仪表盘": [("07-mobile-dashboard.png", "图 2-7　移动端仪表盘（390×844）", MOBILE_CM)],
    "智能客服页": [("08-chatbot.png", "图 2-8　智能客服", DESKTOP_CM)],
    "反爬设置 / 日志页": [("09-anticrawl.png", "图 2-9　反爬与登录态管理", DESKTOP_CM)],
    "系统维护页": [("10b-db-admin.png", "图 2-10　数据库维护", DESKTOP_CM)],
}

PLACEHOLDER = "【此处插入功能截图："
CAPTION_NOTE = "（实际界面示意，敏感信息已打码）"

_OPT_DIR = None


def _embedded_copy(path: str) -> str:
    """图片按原尺寸量化成 256 色后嵌入，docx 体积从 1.5MB 降到 0.8MB。

    界面截图以纯色块和文字为主，256 色看不出差别；原始 PNG 仍留在
    docs/assets/product-shots/ 供后续替换。
    """
    global _OPT_DIR
    from PIL import Image

    if _OPT_DIR is None:
        _OPT_DIR = tempfile.mkdtemp(prefix="xh-doc-img-")
        atexit.register(shutil.rmtree, _OPT_DIR, True)
    out = os.path.join(_OPT_DIR, os.path.basename(path))
    img = Image.open(path).convert("RGB")
    img.quantize(colors=256, method=Image.Quantize.MEDIANCUT,
                 dither=Image.Dither.NONE).save(out, "PNG", optimize=True)
    return out


def _strip_box(cell) -> None:
    """去掉单格虚线表的边框、底色与固定行高，让图片直接铺在正文里。"""
    tc_pr = cell._tc.get_or_add_tcPr()
    for tag in ("w:tcBorders", "w:shd"):
        for el in tc_pr.findall(qn(tag)):
            tc_pr.remove(el)
    borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tc_pr.append(borders)
    tr_pr = cell._tc.getparent().get_or_add_trPr()
    for el in tr_pr.findall(qn("w:trHeight")):
        tr_pr.remove(el)
    # 图片必须整块留在同一页，否则 Word 会把图切成两半（图注被挤到下一页）
    tr_pr.append(OxmlElement("w:cantSplit"))


def _caption(cell, text: str):
    p = cell.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text + CAPTION_NOTE)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x6E, 0x77, 0x81)
    return p


def _fill(cell, shots) -> None:
    _strip_box(cell)
    for p in list(cell.paragraphs):
        p._p.getparent().remove(p._p)
    for idx, (fname, caption, width) in enumerate(shots):
        path = os.path.join(SHOTS, fname)
        if not os.path.exists(path):
            raise SystemExit(f"missing screenshot: {path}")
        p = cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0 if idx == 0 else 6)
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(_embedded_copy(path), width=Cm(width))
        _caption(cell, caption)


def _replace_text(par, text: str) -> None:
    for run in par.runs[1:]:
        run._r.getparent().remove(run._r)
    par.runs[0].text = text


def _all_paragraphs(doc):
    yield from doc.paragraphs
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def main() -> None:
    doc = Document(DOCX)
    done = set()
    for table in doc.tables:
        if len(table.rows) != 1 or len(table.columns) != 1:
            continue
        cell = table.rows[0].cells[0]
        text = cell.text or ""
        if PLACEHOLDER not in text:
            continue
        label = text.split(PLACEHOLDER, 1)[1].split("】", 1)[0].strip()
        shots = PLAN.get(label)
        if not shots:
            raise SystemExit(f"no shot mapped for placeholder: {label}")
        _fill(cell, shots)
        done.add(label)

    missing = set(PLAN) - done
    if missing:
        raise SystemExit(f"placeholders not found: {sorted(missing)}")

    for par in _all_paragraphs(doc):
        t = (par.text or "").strip()
        if t.startswith("说明：本文为产品推广材料"):
            _replace_text(
                par,
                "说明：本文为产品推广材料，功能与配置以当前发布版本为准；"
                "文中截图均取自实际运行界面（示例数据），账号昵称、订单号、"
                "Webhook 地址与令牌等敏感信息已做打码处理。",
            )
        elif t.startswith("替换正文虚线框时按下表取图"):
            _replace_text(
                par,
                "正文配图均按下表取自实际界面，统一 1600×900（移动端 390×844），"
                "并按备注完成打码；后续更新界面时照此清单复核替换。",
            )

    for tbl in doc.tables:
        for row in tbl.rows:
            cells = [c.text.strip() for c in row.cells]
            if "移动端仪表盘" in cells and "建议与桌面端并排做双端对比" in cells:
                _replace_text(
                    row.cells[cells.index("建议与桌面端并排做双端对比")].paragraphs[0],
                    "手机竖屏 390×844，与桌面端共用同一后端",
                )

    doc.save(DOCX)
    print("embedded", len(done), "placeholders ->", DOCX)


if __name__ == "__main__":
    main()
